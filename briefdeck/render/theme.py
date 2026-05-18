"""Industrial navy/gold theme — colour constants + reusable layout primitives."""

from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

# ─── Palette ──────────────────────────────────────────────────────────────────
NAVY = RGBColor(0, 51, 102)
NAVY_DARK = RGBColor(10, 40, 80)
GOLD = RGBColor(200, 150, 30)
WHITE = RGBColor(255, 255, 255)
LIGHT_GREY = RGBColor(245, 245, 248)
DARK_GREY = RGBColor(60, 60, 70)
MID_GREY = RGBColor(130, 155, 180)
PALE_BLUE = RGBColor(180, 200, 220)
GREEN = RGBColor(34, 120, 60)
LIGHT_GREEN = RGBColor(230, 248, 235)
RED_DARK = RGBColor(160, 30, 30)
LIGHT_RED = RGBColor(255, 240, 240)
AMBER = RGBColor(200, 130, 0)
LIGHT_AMBER = RGBColor(255, 250, 220)
BLUE = RGBColor(30, 90, 180)
LIGHT_BLUE = RGBColor(235, 242, 255)
VIOLET = RGBColor(120, 80, 180)
LIGHT_VIOLET = RGBColor(240, 235, 250)

# ─── Slide dimensions (16:9) ──────────────────────────────────────────────────
SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# Accent colour rotation for variety
ACCENT_ROTATION = [GOLD, BLUE, GREEN, AMBER, VIOLET, RED_DARK]


def set_bg(slide, color):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = color


def add_text_box(slide, text, left, top, width, height,
                 font_size=14, bold=False, color=None,
                 align=PP_ALIGN.LEFT, italic=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    if color:
        run.font.color.rgb = color
    return box


def add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape


def add_multi_line(slide, lines, left, top, width, height,
                   font_size=12, color=None, line_spacing=4, bold=False):
    box = slide.shapes.add_textbox(left, top, width, height)
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_before = Pt(line_spacing)
        run = p.add_run()
        run.text = line
        run.font.size = Pt(font_size)
        run.font.bold = bold
        if color:
            run.font.color.rgb = color


def add_header_bar(slide, title):
    """Standard NAVY header bar with gold underline."""
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.9), NAVY)
    add_rect(slide, Inches(0), Inches(0.9), Inches(13.33), Inches(0.06), GOLD)
    add_text_box(slide, title, Inches(0.4), Inches(0.1),
                 Inches(12.5), Inches(0.7),
                 font_size=24, bold=True, color=WHITE)


def add_gold_borders(slide):
    """Top + bottom gold strips — used on dark-background slides."""
    add_rect(slide, Inches(0), Inches(0), Inches(13.33), Inches(0.08), GOLD)
    add_rect(slide, Inches(0), Inches(7.42), Inches(13.33), Inches(0.08), GOLD)


def truncate(text, max_len=120):
    """Defensive truncation for bullet strings that might exceed slide width."""
    if not text:
        return ""
    text = str(text).strip()
    if len(text) <= max_len:
        return text
    return text[: max_len - 1].rstrip() + "…"
